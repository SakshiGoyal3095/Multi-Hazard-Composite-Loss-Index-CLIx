import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from scipy import stats
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def check_distribution(df, dim_map, skew_threshold=2.0, kurtosis_threshold=3.5):
    indicator_cols = [c for c in df.columns if c not in ('District', 'Block')]
    rows = []
    for col in indicator_cols:
        data = df[col]
        skew = stats.skew(data, bias=False)
        kurt = stats.kurtosis(data, bias=False, fisher=False)
        rows.append({
            'Indicator': col,
            'Dimension': dim_map[col],
            'Min': data.min(),
            'Max': data.max(),
            'Mean': data.mean(),
            'Median': data.median(),
            'Skewness': skew,
            'Direction': 'Right' if skew > 0 else 'Left' if skew < 0 else 'None',
            'Kurtosis': kurt,
            'Exceeds_Skew_Threshold': abs(skew) > skew_threshold,
            'Exceeds_Kurtosis_Threshold': kurt > kurtosis_threshold
        })
    dist_df = pd.DataFrame(rows).sort_values(['Dimension', 'Indicator']).reset_index(drop=True)
    return dist_df

# Sibilia et al. (2024): winsorise when skew and kurtosis both exceed thresholds,
# then ln(x+1) only if skew still exceeds log_threshold
def apply_skew_kurtosis_treatment(df, dim_map, skew_threshold=2.0, kurtosis_threshold=3.5,
                                  log_threshold=3.5, winsor_lower=0.02, winsor_upper=0.98):
    treated_df = df.copy()
    treatment_log = []

    indicator_cols = [c for c in df.columns if c not in ('District', 'Block')]
    for col in indicator_cols:
        raw_skew = stats.skew(df[col], bias=False)
        raw_kurt = stats.kurtosis(df[col], bias=False, fisher=False)

        winsorized = False
        log_transformed = False
        current_values = df[col].copy()

        if abs(raw_skew) > skew_threshold and raw_kurt > kurtosis_threshold:
            lower_bound = current_values.quantile(winsor_lower)
            upper_bound = current_values.quantile(winsor_upper)
            current_values = current_values.clip(lower=lower_bound, upper=upper_bound)
            winsorized = True

        post_skew = stats.skew(current_values, bias=False)
        if abs(post_skew) > log_threshold:
            current_values = np.log(current_values + 1)
            log_transformed = True

        treated_df[col] = current_values
        treatment_log.append({
            'Indicator': col,
            'Dimension': dim_map[col],
            'Raw_Skewness': raw_skew,
            'Raw_Kurtosis': raw_kurt,
            'Winsorized': winsorized,
            'Post_Winsorize_Skewness': post_skew if winsorized else raw_skew,
            'Log_Transformed': log_transformed,
            'Final_Skewness': stats.skew(current_values, bias=False)
        })

    treatment_log_df = pd.DataFrame(treatment_log).sort_values(['Dimension', 'Indicator']).reset_index(drop=True)
    return treated_df, treatment_log_df

# winsorising sparse indicators can collapse them to a single value
def drop_post_treatment_zero_variance(df):
    zero_var_cols = [col for col in df.columns if col not in ('District', 'Block') and df[col].std() == 0]
    cleaned = df.drop(columns=zero_var_cols)
    return cleaned, zero_var_cols

def build_treatment_summary(treatment_log_df, log_threshold=3.5):
    still_exceeding = treatment_log_df[treatment_log_df['Final_Skewness'].abs() > log_threshold]
    summary_rows = [{
        'Total Indicators': len(treatment_log_df),
        'Winsorized': treatment_log_df['Winsorized'].sum(),
        'Log-Transformed': treatment_log_df['Log_Transformed'].sum(),
        'No Treatment Needed': (~treatment_log_df['Winsorized'] & ~treatment_log_df['Log_Transformed']).sum(),
        f'Still Exceeding |Skew| > {log_threshold} After Treatment': len(still_exceeding),
        'Indicators Still Exceeding': ', '.join(still_exceeding['Indicator'].tolist()) if len(still_exceeding) else 'None',
        'Method': 'Sibilia et al. (2024): winsorise when skew and kurtosis exceed thresholds; ln(x+1) if skew still exceeds log threshold.'
    }]
    return pd.DataFrame(summary_rows)

def export_distribution_step(dist_df, output_path):
    with pd.ExcelWriter(output_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        dist_df.to_excel(writer, sheet_name='Distribution_Check', index=False)

    wb = load_workbook(output_path)
    orange_fill = PatternFill(start_color='FFD9A0', end_color='FFD9A0', fill_type='solid')
    ws = wb['Distribution_Check']
    header = [cell.value for cell in ws[1]]
    skew_col = header.index('Exceeds_Skew_Threshold') + 1
    kurt_col = header.index('Exceeds_Kurtosis_Threshold') + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=skew_col).value == True or ws.cell(row=row, column=kurt_col).value == True:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = orange_fill
    wb.save(output_path)

def export_treatment_step(treatment_log_df, treated_df, summary_df, output_path, log_threshold=3.5):
    with pd.ExcelWriter(output_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        treatment_log_df.to_excel(writer, sheet_name='Skew_Kurtosis_Treatment', index=False)
        treated_df.to_excel(writer, sheet_name='Treated_Data', index=False)
        summary_df.to_excel(writer, sheet_name='Treatment_Summary', index=False)

    wb = load_workbook(output_path)
    yellow_fill = PatternFill(start_color='FFF2A6', end_color='FFF2A6', fill_type='solid')
    ws = wb['Skew_Kurtosis_Treatment']
    header = [cell.value for cell in ws[1]]
    final_skew_col = header.index('Final_Skewness') + 1
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row=row, column=final_skew_col).value
        if val is not None and abs(val) > log_threshold:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = yellow_fill
    wb.save(output_path)

# Plotting figures 
# one histogram per indicator, before any treatment
def save_raw_histograms(df, output_folder):
    os.makedirs(output_folder, exist_ok=True)
    indicator_cols = [c for c in df.columns if c not in ('District', 'Block')]
    for col in indicator_cols:
        data = df[col]
        fig, ax = plt.subplots(figsize=(5, 4))
        ax.hist(data, bins=15, edgecolor='black', color='skyblue')
        ax.axvline(data.mean(), color='red', linestyle='--', linewidth=1.2, label=f'Mean = {data.mean():.2f}')
        ax.axvline(data.median(), color='blue', linestyle='--', linewidth=1.2, label=f'Median = {data.median():.2f}')
        ax.set_title(col, fontsize=9)
        ax.set_xlabel(col)
        ax.set_ylabel('Frequency')
        ax.legend(fontsize=7)
        stats_text = f"Min = {data.min():.2f}\nMax = {data.max():.2f}\nSkew = {stats.skew(data, bias=False):+.2f}"
        ax.text(0.98, 0.75, stats_text, transform=ax.transAxes, fontsize=7, verticalalignment='top',
                horizontalalignment='right', bbox=dict(boxstyle='round', facecolor='white', alpha=0.8))
        plt.tight_layout()
        fig.savefig(os.path.join(output_folder, f'{col}.png'), dpi=150)
        plt.close(fig)


# one histogram per indicator after treatment, labelled with the treatment applied
def save_treated_histograms(df, treatment_log_df, output_folder):
    os.makedirs(output_folder, exist_ok=True)
    stats_lookup = treatment_log_df.set_index('Indicator')
    indicator_cols = [c for c in df.columns if c not in ('District', 'Block')]
    for col in indicator_cols:
        data = df[col]
        treatment = [name for name, applied in [('Winsorized', stats_lookup.loc[col, 'Winsorized']),
                                                ('Log-Transformed', stats_lookup.loc[col, 'Log_Transformed'])] if applied]
        fig, ax = plt.subplots(figsize=(5, 4))
        ax.hist(data, bins=15, edgecolor='black', color='seagreen')
        ax.axvline(data.mean(), color='red', linestyle='--', linewidth=1.2, label=f'Mean = {data.mean():.2f}')
        ax.axvline(data.median(), color='blue', linestyle='--', linewidth=1.2, label=f'Median = {data.median():.2f}')
        ax.set_title(col, fontsize=9)
        ax.set_xlabel(col)
        ax.set_ylabel('Frequency')
        ax.legend(fontsize=7)
        stats_text = (f"Min = {data.min():.2f}\nMax = {data.max():.2f}\n"
                      f"Skew = {stats_lookup.loc[col, 'Final_Skewness']:+.2f}\n"
                      f"{' + '.join(treatment) if treatment else 'No treatment'}")
        ax.text(0.98, 0.75, stats_text, transform=ax.transAxes, fontsize=7, verticalalignment='top',
                horizontalalignment='right', bbox=dict(boxstyle='round', facecolor='white', alpha=0.8))
        plt.tight_layout()
        fig.savefig(os.path.join(output_folder, f'{col}.png'), dpi=150)
        plt.close(fig)

        
## Note on residual skewness:
# Indicators that remain above the skewness threshold after winsorisation and log-transformation are retained.
# Their influence on the PCA is limited because component extraction uses Spearman rank correlations, which are robust to skewed distributions (Tabachnick & Fidell, 2013).  