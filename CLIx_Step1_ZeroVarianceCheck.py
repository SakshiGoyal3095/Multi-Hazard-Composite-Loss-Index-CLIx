import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def check_zero_variance(df, dim_map):
    indicator_cols = [c for c in df.columns if c not in ('District', 'Block')]
    rows = []
    for col in indicator_cols:
        var = df[col].var()
        rows.append({
            'Indicator': col,
            'Dimension': dim_map[col],
            'Variance': var,
            'Zero_Variance': var == 0
        })
    check_df = pd.DataFrame(rows).sort_values(['Dimension', 'Indicator']).reset_index(drop=True)
    return check_df

def drop_zero_variance(df, check_df):
    zero_var_indicators = check_df.loc[check_df['Zero_Variance'], 'Indicator'].tolist()
    cleaned = df.drop(columns=zero_var_indicators)
    return cleaned, zero_var_indicators

def build_summary(dropped_indicators, total_before, n_blocks):
    total_after = total_before - len(dropped_indicators)
    summary_rows = [{
        'Indicators Before Exclusion': total_before,
        'Zero-Variance Indicators Dropped': len(dropped_indicators),
        'Dropped Indicator Names': ', '.join(dropped_indicators) if dropped_indicators else 'None',
        'Indicators Remaining': total_after,
        'Reason for Exclusion': f'Indicator recorded an identical value across all {n_blocks} blocks, providing no variance for PCA (Jolliffe, 2002).'
    }]
    return pd.DataFrame(summary_rows)

# output_path is the results workbook, never the raw input file
def export_variance_step(check_df, cleaned_df, summary_df, output_path):
    with pd.ExcelWriter(output_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        check_df.to_excel(writer, sheet_name='Variance_Check', index=False)
        cleaned_df.to_excel(writer, sheet_name='Cleaned_Data', index=False)
        summary_df.to_excel(writer, sheet_name='Summary', index=False)

    wb = load_workbook(output_path)
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    ws = wb['Variance_Check']
    header = [cell.value for cell in ws[1]]
    zero_var_col = header.index('Zero_Variance') + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=zero_var_col).value == True:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = red_fill
    wb.save(output_path)